"""
Exports a list of per-PDF report dicts (see report.py) to Markdown, HTML,
or XLSX. Results stay grouped by PDF in every format.
"""
import html as _html
import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

STATUS_COLORS = {
    "pass": "1E7B34",
    "fail": "C0392B",
    "na": "7F8C8D",
    "manual": "B7791F",
    "review": "8E44AD",
}
STATUS_FILL = {
    "pass": "E9F7EF",
    "fail": "FDEDEC",
    "na": "F2F3F4",
    "manual": "FEF5E7",
    "review": "F4ECF7",
}


def _overall_totals(reports):
    """Combined pass/fail/etc counts and conformance score across every PDF
    in the batch."""
    totals = {"pass": 0, "fail": 0, "na": 0, "manual": 0, "review": 0}
    for r in reports:
        s = r["summary"]
        for k in totals:
            totals[k] += s.get(k, 0)
    scored = totals["pass"] + totals["fail"]
    totals["score_pct"] = round(100 * totals["pass"] / scored) if scored else None
    return totals


def _pages_str(row):
    pages = row.get("pages") or []
    return ", ".join(str(p) for p in pages) if pages else "-"


# ============================================================ Markdown =====
def to_markdown(reports, title="PDF Accessibility Report"):
    lines = [f"# {title}", "", f"_Generated {datetime.datetime.now():%Y-%m-%d %H:%M}_", ""]
    if len(reports) > 1:
        overall = _overall_totals(reports)
        overall_score = f"{overall['score_pct']}%" if overall["score_pct"] is not None else "n/a"
        lines.append(f"**{len(reports)} PDF(s) tested. Overall conformance score: {overall_score}.**")
        lines.append("")
        lines.append("| # | File | Pass | Fail | N/A | Review | Manual | Score |")
        lines.append("|---|------|------|------|-----|--------|--------|-------|")
        for i, r in enumerate(reports, start=1):
            s = r["summary"]
            score = f"{s['score_pct']}%" if s["score_pct"] is not None else "-"
            lines.append(f"| {i} | {r['filename']} | {s['pass']} | {s['fail']} | "
                          f"{s['na']} | {s['review']} | {s['manual']} | {score} |")
        lines.append("")

    for r in reports:
        lines.append(f"## {r['filename']}")
        lines.append("")
        if r.get("source_url"):
            lines.append(f"- **Source URL:** {r['source_url']}")
        if r.get("source_page"):
            lines.append(f"- **Found on page:** {r['source_page']}")
        if r.get("page_count") is not None:
            lines.append(f"- **Pages:** {r['page_count']}")
        if r.get("title"):
            lines.append(f"- **Document title:** {r['title']}")
        lines.append(f"- **Tagged PDF:** {'Yes' if r.get('tagged') else 'No'}")
        s = r["summary"]
        lines.append(f"- **Results:** {s['pass']} Pass / {s['fail']} Fail / "
                      f"{s['na']} N/A / {s['review']} Needs review / "
                      f"{s['manual']} Manual check needed")
        if r.get("errors"):
            lines.append(f"- **Parser notes:** {'; '.join(r['errors'])}")
        lines.append("")

        for row in r["rows"]:
            wcag_str = "; ".join(f"{w['criterion']} ({w['level']})" for w in row["wcag"]) or "-"
            lines.append(f"### {row['id']} - {row['title']}  [{row['status_label']}]")
            lines.append("")
            lines.append(f"- **Category:** {row['category']}")
            lines.append(f"- **WCAG:** {wcag_str}")
            lines.append(f"- **PDF/UA:** {row['pdfua']}")
            lines.append(f"- **Finding:** {row['summary']}")
            if row.get("pages"):
                lines.append(f"- **Location:** Page(s) {_pages_str(row)}")
            if row["evidence"]:
                lines.append("- **Details:**")
                for e in row["evidence"]:
                    lines.append(f"  - {e}")
            if row["status"] in ("fail", "manual", "review") and row["remediation"]:
                label = "How to verify manually" if row["status"] == "manual" else "Remediation steps"
                lines.append(f"- **{label}:**")
                for i, step in enumerate(row["remediation"], start=1):
                    lines.append(f"  {i}. {step}")
            lines.append("")
    return "\n".join(lines)


# ================================================================ HTML =====
def to_html(reports, title="PDF Accessibility Report"):
    def esc(s):
        return _html.escape(str(s))

    css = """
    body{font-family:-apple-system,Segoe UI,Roboto,Helvetica,Arial,sans-serif;
         max-width:1100px;margin:0 auto;padding:24px;color:#1b2430;background:#fafbfc;}
    h1{margin-bottom:4px;} .meta{color:#667;font-size:14px;margin-bottom:24px;}
    table{border-collapse:collapse;width:100%;margin-bottom:24px;background:#fff;}
    th,td{border:1px solid #dde2e8;padding:8px 10px;font-size:14px;text-align:left;}
    th{background:#eef1f5;}
    .pdfblock{border:1px solid #dde2e8;border-radius:10px;padding:20px 24px;
              margin-bottom:32px;background:#fff;}
    .pdfhead{display:flex;justify-content:space-between;align-items:baseline;
              flex-wrap:wrap;gap:8px;}
    .badge{display:inline-block;padding:2px 10px;border-radius:12px;font-size:12px;
           font-weight:600;color:#fff;}
    .row{border:1px solid #e3e7ec;border-radius:8px;padding:12px 16px;margin:10px 0;}
    .row h4{margin:0 0 6px 0;font-size:15px;}
    .rowmeta{font-size:12px;color:#556;margin-bottom:6px;}
    .finding{font-size:14px;margin:6px 0;}
    ul{margin:6px 0;padding-left:20px;font-size:13px;}
    .steps{background:#f7f8fa;border-radius:6px;padding:8px 14px;margin-top:6px;}
    .summarybar{display:flex;gap:10px;flex-wrap:wrap;margin:8px 0 16px 0;}
    .chip{border-radius:8px;padding:6px 12px;font-size:13px;font-weight:600;}
    """

    parts = [f"<!DOCTYPE html><html><head><meta charset='utf-8'>"
             f"<title>{esc(title)}</title><style>{css}</style></head><body>"]
    parts.append(f"<h1>{esc(title)}</h1>")
    overall = _overall_totals(reports)
    overall_score = f"{overall['score_pct']}%" if overall["score_pct"] is not None else "n/a"
    parts.append(f"<div class='meta'>Generated {datetime.datetime.now():%Y-%m-%d %H:%M} "
                 f"&middot; {len(reports)} PDF(s) tested "
                 f"&middot; Overall conformance score: {esc(overall_score)}</div>")

    if len(reports) > 1:
        parts.append("<table><tr><th>#</th><th>File</th><th>Pass</th><th>Fail</th>"
                      "<th>N/A</th><th>Review</th><th>Manual</th><th>Score</th></tr>")
        for i, r in enumerate(reports, start=1):
            s = r["summary"]
            score = f"{s['score_pct']}%" if s["score_pct"] is not None else "-"
            parts.append(f"<tr><td>{i}</td><td>{esc(r['filename'])}</td>"
                         f"<td>{s['pass']}</td><td>{s['fail']}</td><td>{s['na']}</td>"
                         f"<td>{s['review']}</td><td>{s['manual']}</td><td>{score}</td></tr>")
        parts.append("</table>")

    for r in reports:
        s = r["summary"]
        score = f"{s['score_pct']}%" if s["score_pct"] is not None else "n/a"
        parts.append("<div class='pdfblock'>")
        parts.append(f"<div class='pdfhead'><h2>{esc(r['filename'])}</h2>"
                     f"<span class='badge' style='background:#2c3e50'>Score: {score}</span></div>")
        meta_bits = []
        if r.get("source_url"):
            meta_bits.append(f"Source: {esc(r['source_url'])}")
        if r.get("page_count") is not None:
            meta_bits.append(f"{r['page_count']} pages")
        meta_bits.append(f"Tagged PDF: {'Yes' if r.get('tagged') else 'No'}")
        parts.append(f"<div class='meta'>{' &middot; '.join(meta_bits)}</div>")
        parts.append("<div class='summarybar'>"
                     f"<span class='chip' style='background:{STATUS_FILL['pass']};color:{STATUS_COLORS['pass']}'>Pass {s['pass']}</span>"
                     f"<span class='chip' style='background:{STATUS_FILL['fail']};color:{STATUS_COLORS['fail']}'>Fail {s['fail']}</span>"
                     f"<span class='chip' style='background:{STATUS_FILL['review']};color:{STATUS_COLORS['review']}'>Needs review {s['review']}</span>"
                     f"<span class='chip' style='background:{STATUS_FILL['manual']};color:{STATUS_COLORS['manual']}'>Manual {s['manual']}</span>"
                     f"<span class='chip' style='background:{STATUS_FILL['na']};color:{STATUS_COLORS['na']}'>N/A {s['na']}</span>"
                     "</div>")
        for row in r["rows"]:
            color = STATUS_COLORS.get(row["status"], "555")
            fill = STATUS_FILL.get(row["status"], "#eee")
            wcag_str = "; ".join(f"{esc(w['criterion'])} ({esc(w['level'])})" for w in row["wcag"]) or "-"
            parts.append("<div class='row'>")
            parts.append(f"<h4>{esc(row['id'])} &mdash; {esc(row['title'])} "
                         f"<span class='badge' style='background:{color}'>{esc(row['status_label'])}</span></h4>")
            parts.append(f"<div class='rowmeta'>Category: {esc(row['category'])} &middot; "
                         f"WCAG: {wcag_str} &middot; PDF/UA: {esc(row['pdfua'])}</div>")
            parts.append(f"<div class='finding'>{esc(row['summary'])}</div>")
            if row.get("pages"):
                parts.append(f"<div class='finding'><strong>Location:</strong> "
                             f"Page(s) {esc(_pages_str(row))}</div>")
            if row["evidence"]:
                parts.append("<ul>" + "".join(f"<li>{esc(e)}</li>" for e in row["evidence"]) + "</ul>")
            if row["status"] in ("fail", "manual", "review") and row["remediation"]:
                label = "How to verify manually" if row["status"] == "manual" else "Remediation steps"
                parts.append(f"<div class='steps'><strong>{label}:</strong><ol>" +
                             "".join(f"<li>{esc(step)}</li>" for step in row["remediation"]) +
                             "</ol></div>")
            parts.append("</div>")
        parts.append("</div>")

    parts.append("</body></html>")
    return "\n".join(parts)


# ================================================================ XLSX =====
LINK_FONT = Font(color="0563C1", underline="single")


def to_xlsx(reports, path, title="PDF Accessibility Report"):
    wb = Workbook()

    # Decide each PDF's sheet name up front (deduped) so the Summary sheet
    # can link straight to it before the sheet itself is created.
    used_names = set()
    sheet_names = []
    for r in reports:
        base = _safe_sheet_name(r["filename"])
        name = base
        n = 2
        while name in used_names:
            name = _safe_sheet_name(f"{base[:28]}_{n}")
            n += 1
        used_names.add(name)
        sheet_names.append(name)

    overall = _overall_totals(reports)
    overall_score = f"{overall['score_pct']}%" if overall["score_pct"] is not None else "n/a"

    # ---- Summary sheet ----
    ws = wb.active
    ws.title = "Summary"
    ws.append([title])
    ws["A1"].font = Font(size=16, bold=True)
    ws.append([f"Generated {datetime.datetime.now():%Y-%m-%d %H:%M}"])
    ws.append([f"{len(reports)} PDF(s) tested - "
               f"Overall conformance score (pass / (pass+fail) across the batch): {overall_score}"])
    ws["A3"].font = Font(size=12, bold=True, color="2C3E50")
    ws.append([])
    header_row = 5
    header = ["#", "PDF title (click to open its results tab)", "Pages", "Tagged",
               "Pass", "Fail", "N/A", "Needs review", "Manual", "Score"]
    ws.append(header)
    for c in ws[header_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2C3E50")
    for i, (r, sheet_name) in enumerate(zip(reports, sheet_names), start=1):
        s = r["summary"]
        ws.append([i, r["filename"], r.get("page_count"),
                   "Yes" if r.get("tagged") else "No",
                   s["pass"], s["fail"], s["na"], s["review"], s["manual"],
                   f"{s['score_pct']}%" if s["score_pct"] is not None else "n/a"])
        cell = ws.cell(row=ws.max_row, column=2)
        cell.hyperlink = f"#'{sheet_name}'!A1"
        cell.font = LINK_FONT
    total_row_idx = ws.max_row + 1
    ws.append(["", "TOTAL (all PDFs in this batch)", "", "",
               overall["pass"], overall["fail"], overall["na"],
               overall["review"], overall["manual"], overall_score])
    for c in ws[total_row_idx]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EEF1F5")
    for col, width in zip("ABCDEFGHIJ", [4, 46, 8, 8, 7, 7, 6, 12, 8, 8]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = f"A{header_row + 1}"

    # ---- One sheet per PDF (grouped findings) ----
    for r, sheet_name in zip(reports, sheet_names):
        ws2 = wb.create_sheet(sheet_name)
        ws2.append(["<< Back to Summary"])
        back_cell = ws2.cell(row=1, column=1)
        back_cell.hyperlink = "#'Summary'!A1"
        back_cell.font = LINK_FONT
        s = r["summary"]
        score = f"{s['score_pct']}%" if s["score_pct"] is not None else "n/a"
        ws2.append([r["filename"], f"Conformance score: {score}"])
        ws2["A2"].font = Font(size=13, bold=True)
        ws2.append([])
        cols = ["Test ID", "Title", "Category", "Status", "Finding",
                "Page(s)", "Details", "WCAG", "PDF/UA", "Automation",
                "Remediation / Manual steps"]
        ws2.append(cols)
        header_row_idx = ws2.max_row
        for c in ws2[header_row_idx]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2C3E50")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for row in r["rows"]:
            wcag_str = "; ".join(f"{w['criterion']} ({w['level']})" for w in row["wcag"]) or "-"
            details = "\n".join(row["evidence"])
            remediation = "\n".join(
                f"{i}. {s}" for i, s in enumerate(row["remediation"], start=1))
            ws2.append([row["id"], row["title"], row["category"],
                        row["status_label"], row["summary"], _pages_str(row),
                        details, wcag_str, row["pdfua"], row["automation"],
                        remediation])
            fill = STATUS_FILL.get(row["status"], "FFFFFF")
            last_row = ws2.max_row
            for c in ws2[last_row]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            ws2.cell(row=last_row, column=4).fill = PatternFill("solid", fgColor=fill)
        widths = [10, 26, 14, 14, 34, 10, 34, 22, 20, 10, 46]
        for i, w in enumerate(widths, start=1):
            ws2.column_dimensions[get_column_letter(i)].width = w
        ws2.freeze_panes = f"A{header_row_idx + 1}"

    wb.save(path)
    return path


def _safe_sheet_name(name):
    bad = set('[]:*?/\\')
    cleaned = "".join(c for c in name if c not in bad)
    return cleaned[:31] or "PDF"
